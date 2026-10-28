"""
Document loader module supporting PDF, TXT, and DOCX formats.
Includes validation, cleaning, metadata preservation, and duplicate detection hashing.
"""

from __future__ import annotations

import io
import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from config import settings
from src.utils import calculate_file_hash, clean_text, get_logger

logger = get_logger(__name__)


class DocumentLoaderError(Exception):
    """Base exception for document loader errors."""
    pass


class UnsupportedFileTypeError(DocumentLoaderError):
    """Raised when an unsupported file extension is uploaded."""
    pass


class FileTooLargeError(DocumentLoaderError):
    """Raised when an uploaded file exceeds the configured size limit."""
    pass


class CorruptedFileError(DocumentLoaderError):
    """Raised when a document cannot be parsed or read."""
    pass


class EmptyDocumentError(DocumentLoaderError):
    """Raised when an uploaded document contains no readable text."""
    pass


@dataclass
class LoadedDocument:
    """Represents an extracted document page or section with metadata."""
    page_content: str
    metadata: Dict[str, Any] = field(default_factory=dict)

    def to_langchain_doc(self):
        """Converts to a LangChain Document if langchain_core is present."""
        try:
            from langchain_core.documents import Document
            return Document(page_content=self.page_content, metadata=self.metadata)
        except ImportError:
            return self


class DocumentLoader:
    """
    Robust document loader supporting PDF, TXT, and DOCX file processing
    from filesystem paths or in-memory file buffers.
    """

    def __init__(
        self,
        allowed_extensions: Optional[List[str]] = None,
        max_upload_size_mb: Optional[int] = None,
    ):
        self.allowed_extensions = (
            allowed_extensions
            if allowed_extensions is not None
            else settings.allowed_extensions
        )
        self.max_upload_size_mb = (
            max_upload_size_mb
            if max_upload_size_mb is not None
            else settings.max_upload_size_mb
        )

    def validate_file(self, filename: str, file_size: int) -> None:
        """
        Validates file extension and size against configured boundaries.
        
        Args:
            filename: Name of the file with extension.
            file_size: Size in bytes.
        """
        ext = Path(filename).suffix.lower()
        if ext not in self.allowed_extensions:
            raise UnsupportedFileTypeError(
                f"File '{filename}' with extension '{ext}' is not supported. "
                f"Allowed formats: {', '.join(self.allowed_extensions)}"
            )

        max_bytes = self.max_upload_size_mb * 1024 * 1024
        if file_size > max_bytes:
            raise FileTooLargeError(
                f"File '{filename}' ({file_size / (1024 * 1024):.2f} MB) exceeds "
                f"the maximum allowed size of {self.max_upload_size_mb} MB."
            )

    def load_from_path(self, file_path: Union[str, Path]) -> List[LoadedDocument]:
        """Loads and extracts content from a filesystem path."""
        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        file_size = path.stat().st_size
        self.validate_file(path.name, file_size)

        with open(path, "rb") as f:
            content_bytes = f.read()

        return self._process_bytes(
            content_bytes=content_bytes,
            filename=path.name,
            file_size=file_size,
        )

    def load_from_bytes(
        self,
        content_bytes: bytes,
        filename: str,
    ) -> List[LoadedDocument]:
        """Loads and extracts content from raw in-memory bytes (e.g. Streamlit UploadedFile)."""
        file_size = len(content_bytes)
        self.validate_file(filename, file_size)
        return self._process_bytes(
            content_bytes=content_bytes,
            filename=filename,
            file_size=file_size,
        )

    def _process_bytes(
        self,
        content_bytes: bytes,
        filename: str,
        file_size: int,
    ) -> List[LoadedDocument]:
        """Internal dispatch based on file extension."""
        ext = Path(filename).suffix.lower()
        file_hash = calculate_file_hash(content_bytes)

        logger.info("Processing document: %s (size: %d bytes, hash: %s...)", filename, file_size, file_hash[:8])

        if ext == ".pdf":
            docs = self._extract_pdf(content_bytes, filename, file_hash, file_size)
        elif ext == ".docx":
            docs = self._extract_docx(content_bytes, filename, file_hash, file_size)
        elif ext == ".txt":
            docs = self._extract_txt(content_bytes, filename, file_hash, file_size)
        else:
            raise UnsupportedFileTypeError(f"Unsupported extension: {ext}")

        # Check total extracted text
        total_text = "".join(doc.page_content for doc in docs).strip()
        if not total_text:
            raise EmptyDocumentError(
                f"Document '{filename}' contains no readable text or is empty."
            )

        logger.info(
            "Successfully extracted %d section(s)/page(s) from '%s' (%d total characters)",
            len(docs),
            filename,
            len(total_text),
        )
        return docs

    def _extract_pdf(
        self,
        content_bytes: bytes,
        filename: str,
        file_hash: str,
        file_size: int,
    ) -> List[LoadedDocument]:
        """Extracts text page-by-page from a PDF using pypdf."""
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ImportError("pypdf is required to process PDF files. Please install pypdf>=4.3.0.")

        try:
            reader = PdfReader(io.BytesIO(content_bytes))
            if reader.is_encrypted:
                try:
                    # Attempt empty password decryption
                    reader.decrypt("")
                except Exception as e:
                    raise CorruptedFileError(
                        f"PDF '{filename}' is password protected and cannot be read: {e}"
                    )

            docs: List[LoadedDocument] = []
            for page_idx, page in enumerate(reader.pages, start=1):
                try:
                    page_text = page.extract_text() or ""
                    cleaned = clean_text(page_text)
                    if cleaned:
                        metadata = {
                            "source_filename": filename,
                            "file_type": ".pdf",
                            "file_size": file_size,
                            "file_hash": file_hash,
                            "page_number": page_idx,
                            "total_pages": len(reader.pages),
                            "character_count": len(cleaned),
                        }
                        docs.append(LoadedDocument(page_content=cleaned, metadata=metadata))
                except Exception as page_err:
                    logger.warning("Error reading page %d of '%s': %s", page_idx, filename, page_err)

            return docs
        except CorruptedFileError:
            raise
        except Exception as e:
            raise CorruptedFileError(f"Failed to parse PDF document '{filename}': {e}") from e

    def _extract_docx(
        self,
        content_bytes: bytes,
        filename: str,
        file_hash: str,
        file_size: int,
    ) -> List[LoadedDocument]:
        """Extracts text and table content from a Word document using python-docx."""
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise ImportError(
                "python-docx is required to process DOCX files. Please install python-docx>=1.1.2."
            )

        try:
            doc = DocxDocument(io.BytesIO(content_bytes))
            sections: List[str] = []

            # Extract body paragraphs
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    sections.append(p_text)

            # Extract table cells
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_cells:
                        sections.append(" | ".join(row_cells))

            full_text = clean_text("\n\n".join(sections))
            if not full_text:
                return []

            metadata = {
                "source_filename": filename,
                "file_type": ".docx",
                "file_size": file_size,
                "file_hash": file_hash,
                "page_number": 1,
                "character_count": len(full_text),
            }
            return [LoadedDocument(page_content=full_text, metadata=metadata)]
        except Exception as e:
            raise CorruptedFileError(f"Failed to parse DOCX document '{filename}': {e}") from e

    def _extract_txt(
        self,
        content_bytes: bytes,
        filename: str,
        file_hash: str,
        file_size: int,
    ) -> List[LoadedDocument]:
        """Extracts text from plain text file with multi-encoding fallback."""
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"]
        decoded_text: Optional[str] = None

        for enc in encodings:
            try:
                decoded_text = content_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue

        if decoded_text is None:
            # Fallback with replacement
            decoded_text = content_bytes.decode("utf-8", errors="replace")

        cleaned = clean_text(decoded_text)
        if not cleaned:
            return []

        metadata = {
            "source_filename": filename,
            "file_type": ".txt",
            "file_size": file_size,
            "file_hash": file_hash,
            "page_number": 1,
            "character_count": len(cleaned),
        }
        return [LoadedDocument(page_content=cleaned, metadata=metadata)]
